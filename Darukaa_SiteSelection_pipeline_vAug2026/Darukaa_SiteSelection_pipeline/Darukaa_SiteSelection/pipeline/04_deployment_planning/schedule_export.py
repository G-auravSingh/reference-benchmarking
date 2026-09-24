"""
04_deployment_planning / schedule_export.py
==============================================

Exports the deployment schedule as a real CSV/Excel file with calendar
dates, not just day-offsets, matching the Tata Motors deployment summary
format.

If a project hasn't set `project_start_date` in its config.yaml, this
still exports the CSV/Excel using relative day-offset labels ("Day 1",
"Day 8", ...) rather than refusing to run — a schedule without a firm
start date is still real and useful, it just can't show calendar dates
yet. The moment a start date is set, re-running this stage picks it up.
"""
from __future__ import annotations

import csv
import datetime
import logging
import sys
from pathlib import Path
from typing import Any, Dict, List

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "common"))
import config_schema  # noqa: E402

logger = logging.getLogger(__name__)


def _row_date(start_date: datetime.date | None, day_offset: int) -> str:
    if start_date is None:
        return f"Day {day_offset + 1}"
    return (start_date + datetime.timedelta(days=day_offset)).strftime("%Y-%m-%d (%a)")


def _load_position_coords(project_dir: Path) -> Dict[str, tuple]:
    """{position_name: (lat, lon)} from both the audiomoth and camera trap
    pools. Position names come from 03b's output; this is the same
    join-by-name pattern
    used everywhere else in the pipeline."""
    import json
    from shapely.geometry import shape
    coords = {}
    for fname in ("position_pool.geojson", "camera_trap_pool.geojson"):
        path = project_dir / "outputs" / "03b_position_scoring" / fname
        if not path.exists():
            continue
        with open(path) as f:
            fc = json.load(f)
        for feat in fc["features"]:
            name = feat["properties"].get("name")
            if name:
                c = shape(feat["geometry"]).centroid
                coords[name] = (round(c.y, 6), round(c.x, 6))
    return coords


def build_schedule_rows(deployment_report: Dict[str, Any], cfg: Dict[str, Any],
                        position_coords: Dict[str, tuple] | None = None) -> List[Dict[str, Any]]:
    start_date = None
    if cfg.get("project_start_date"):
        start_date = datetime.date.fromisoformat(cfg["project_start_date"])
    position_coords = position_coords or {}

    def _lat_lon(name: str) -> tuple:
        return position_coords.get(name, ("\u2014", "\u2014"))

    regime = deployment_report.get("regime")
    device_counters: Dict[str, int] = {}

    def _device_id(prefix: str, emu_id: str) -> str:
        key = f"{prefix}-{emu_id}"
        device_counters[key] = device_counters.get(key, 0) + 1
        return f"{prefix}-{emu_id}-{device_counters[key]:02d}"

    rows = []
    for c in deployment_report.get("schedule", []):
        deploy_date = _row_date(start_date, c["start_day_offset"])
        retrieve_date = _row_date(start_date, c["recording_end_day_offset"])

        # Both continuous_proportional and stratified_single_pass carry
        # real per-position weekly assignments — sequential_cluster is the
        # only regime without them (it schedules whole EMUs per cycle, not
        # individual positions).
        if c.get("position_assignments"):
            for a in c["position_assignments"]:
                for p in a["positions"]:
                    lat, lon = _lat_lon(p["name"])
                    rows.append({
                        "Cycle/Week": c["cycle_number"],
                        "Deployment date": deploy_date,
                        "Retrieval date": retrieve_date,
                        "Stream": "audiomoth",
                        "Device ID": _device_id("AM", a["emu_id"]),
                        "EMU": a["emu_id"],
                        "Position": p["name"],
                        "Latitude": lat,
                        "Longitude": lon,
                        "Pool wrapped this week": "Yes" if p.get("wrapped") else "No",
                    })
            for p in c.get("camera_trap_positions", []):
                lat, lon = _lat_lon(p["name"])
                rows.append({
                    "Cycle/Week": c["cycle_number"],
                    "Deployment date": deploy_date,
                    "Retrieval date": retrieve_date,
                    "Stream": "camera_trap",
                    "Device ID": _device_id("CAM", p.get("emu_id", "site")),
                    "EMU": p.get("emu_id", "\u2014"),
                    "Position": p["name"],
                    "Latitude": lat,
                    "Longitude": lon,
                    "Pool wrapped this week": "Yes" if p.get("wrapped") else "No",
                })
        else:
            for emu_id in c["emu_ids"]:
                rows.append({
                    "Cycle/Week": c["cycle_number"],
                    "Deployment date": deploy_date,
                    "Retrieval date": retrieve_date,
                    "Stream": "audiomoth",
                    "Device ID": _device_id("AM", emu_id),
                    "EMU": emu_id,
                    "Position": "\u2014 (see field_map.html for the specific candidate)",
                    "Latitude": "\u2014", "Longitude": "\u2014",
                    "Pool wrapped this week": "\u2014",
                })
    return rows


def run_schedule_export(project_dir: str | Path) -> Dict[str, Any]:
    project_dir = Path(project_dir)
    cfg = config_schema.load_config(project_dir / "config.yaml")
    dep_path = project_dir / "outputs" / "04_deployment_planning" / "deployment_schedule.json"
    if not dep_path.exists():
        raise FileNotFoundError(f"{dep_path} not found — run 04_deployment_planning first.")
    import json
    with open(dep_path) as f:
        deployment_report = json.load(f)

    position_coords = _load_position_coords(project_dir)
    rows = build_schedule_rows(deployment_report, cfg, position_coords)
    out_dir = project_dir / "outputs" / "04_deployment_planning"

    csv_path = out_dir / "deployment_schedule.csv"
    fieldnames = ["Cycle/Week", "Deployment date", "Retrieval date", "Stream", "Device ID",
                 "EMU", "Position", "Latitude", "Longitude", "Pool wrapped this week"]
    with open(csv_path, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    xlsx_path = None
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Deployment Schedule"
        ws.append(fieldnames)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="15803D")
        for row in rows:
            ws.append([row[k] for k in fieldnames])
        for col in ws.columns:
            max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(40, max_len + 2)
        xlsx_path = out_dir / "deployment_schedule.xlsx"
        wb.save(xlsx_path)
    except ImportError:
        logger.warning("openpyxl not available — wrote CSV only, no .xlsx.")

    logger.info("Deployment schedule exported: %s (%d rows)%s",
                csv_path, len(rows), f" + {xlsx_path.name}" if xlsx_path else "")
    return {"csv_path": str(csv_path), "xlsx_path": str(xlsx_path) if xlsx_path else None, "n_rows": len(rows)}


if __name__ == "__main__":
    import argparse
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    parser = argparse.ArgumentParser()
    parser.add_argument("project_dir")
    args = parser.parse_args()
    run_schedule_export(args.project_dir)
